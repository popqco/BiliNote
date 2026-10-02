import importlib.util
import pathlib
import shutil
import tempfile
import unittest
import zipfile

ROOT = pathlib.Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / "app" / "utils" / "export.py"
spec = importlib.util.spec_from_file_location("export", MODULE_PATH)
if spec is None or spec.loader is None:
    raise ImportError("export module spec not found")
export = importlib.util.module_from_spec(spec)
spec.loader.exec_module(export)

import fitz  # noqa: E402
from docx import Document  # noqa: E402

# 1x1 红色 PNG（足够让渲染器走一遍图片链路）
TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753de"
    "0000000c4944415408d763f8cfc0000003010100eb1f5e3b0000000049454e44ae426082"
)

SAMPLE_MARKDOWN = f"""# 视频笔记：单元测试样例

## 项目概览

这是一个**导出测试**，包含 *斜体*、~~删除线~~、`行内代码` 和行内公式 $E=mc^2$。

## 功能清单

- 无序列表一
- 无序列表二
  1. 有序子项甲
  2. 有序子项乙

## 数据表

| 功能 | 状态 |
|---|---|
| 表格 | 通过 |
| 公式 | 通过 |

> 引用块：图片来自本地文件。

![截图]({(ROOT / "app").as_uri()})

$$\\int_0^1 x^2 \\, dx = \\frac{{1}}{{3}}$$

```python
def hello():
    print("world")
```

[示例链接](https://example.com)
"""


def _make_local_image(tmpdir: str) -> str:
    path = pathlib.Path(tmpdir) / "sample_screenshot.png"
    path.write_bytes(TINY_PNG)
    return str(path)


class TestExtractMath(unittest.TestCase):
    def test_display_and_inline_math_extracted(self):
        md, items = export.extract_math("前文 $a^2+b^2=c^2$ 后文\n\n$$x = \\frac{1}{2}$$\n\n尾")
        self.assertEqual(len(items), 2)
        self.assertEqual(sum(1 for it in items if it["display"]), 1)
        self.assertEqual(sum(1 for it in items if not it["display"]), 1)
        self.assertNotIn("$", md)
        self.assertIn("\uE0000\uE001", md)
        self.assertIn("\uE0001\uE001", md)

    def test_math_inside_code_fence_untouched(self):
        md = "```python\nx = '$not_math$\n'\n```\n\n正文 $y=2$ 结束"
        _, items = export.extract_math(md)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["latex"], "y=2")

    def test_currency_not_treated_as_math(self):
        _, items = export.extract_math("花 5$ 买一杯，再花 $10 买一包，总计 $15")
        self.assertEqual(items, [])


class TestExportPdf(unittest.TestCase):
    def test_pdf_generated_with_table_and_formula(self):
        tmpdir = tempfile.mkdtemp()
        try:
            img = _make_local_image(tmpdir)
            md = SAMPLE_MARKDOWN.replace((ROOT / "app").as_uri(), img)
            path = export.export_pdf(md, "测试:导出/样例", out_dir=tmpdir)
            self.assertTrue(path.endswith(".pdf"))
            doc = fitz.open(path)
            try:
                self.assertGreaterEqual(len(doc), 1)
                text = "".join(page.get_text() for page in doc)
                self.assertIn("视频笔记", text)
                self.assertIn("表格", text)          # 表格渲染成正文而非竖线原文
                self.assertNotIn("|---|", text)
                # 截图 + 两个公式图（行内/独立）都应内嵌为图片
                total_images = sum(len(p.get_images()) for p in doc)
                self.assertGreaterEqual(total_images, 3)
                # 公式不允许静默降级成源码文本（\frac 丢失反斜杠后是 "rac{"）
                self.assertNotIn("rac{", text)
                self.assertNotIn("\\frac", text)
            finally:
                doc.close()
        finally:
            # Windows 下 fitz 句柄释放有延迟，直接忽略清理失败
            shutil.rmtree(tmpdir, ignore_errors=True)


class TestExportDocx(unittest.TestCase):
    def _read_document_xml(self, path: str) -> str:
        with zipfile.ZipFile(path) as zf:
            return zf.read("word/document.xml").decode("utf-8")

    def test_docx_generated_with_formula_and_table(self):
        tmpdir = tempfile.mkdtemp()
        try:
            img = _make_local_image(tmpdir)
            md = SAMPLE_MARKDOWN.replace((ROOT / "app").as_uri(), img)
            path = export.export_docx(md, "测试:导出/样例", out_dir=tmpdir)
            self.assertTrue(path.endswith(".docx"))

            doc = Document(path)
            headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            self.assertTrue(any("单元测试样例" in h for h in headings))
            self.assertGreater(len(doc.tables), 0)          # GFM 表格
            self.assertGreaterEqual(len(doc.inline_shapes), 1)  # 截图（公式走 OMML，不算图片）
            # 行内/独立公式至少产出 OMML（含 oMath 标签）
            xml = self._read_document_xml(path)
            self.assertIn("oMath", xml)
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)

    def test_safe_filename_strips_illegal_chars(self):
        self.assertEqual(export.safe_filename('a/b\\c:d*e?f"g<h>i|j'), "a_b_c_d_e_f_g_h_i_j")
        self.assertEqual(export.safe_filename("  "), "note")


if __name__ == "__main__":
    unittest.main()
