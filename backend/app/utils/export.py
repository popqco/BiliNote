# app/utils/export.py
# 笔记导出：Markdown → PDF / Word(docx)
# 纯工具模块：不 import app 包内模块，tests 按文件路径加载（照 test_note_helper 模式）。
#
# 管线：
#   1) 数学公式预处理：$...$ / $$...$$ 挖成占位符（\uE000{i}\uE001，私有区字符不与正文冲突），
#      代码块 / 行内代码先保护再挖，避免误伤
#   2) PDF：占位符替换为 mathtext 渲出的 PNG data-URI → markdown-it-py 转 HTML
#      → fitz.Story(带 CSS) 分页写 PDF。不走 markdown_pdf：1.7 的 Section 没有
#      user_css 入口且表格样式不可控，Story 直连全部可控（探针实测验证）
#   3) Word：占位符还原为 OMML 原生公式对象（可编辑），markdown-it-py token 流
#      → python-docx 逐元素生成
#
# 图片定位：markdown 里的 /static/... 相对 backend 运行根解析成本地文件；data:/本地绝对路径
# 直接用；http(s) 下载（MuPDF Story 不联网，PDF 侧远程图必须落成 data-URI）。

import base64
import io
import mimetypes
import os
import re
import uuid
from urllib.parse import unquote, urlparse

import httpx

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
EXPORT_DIR = os.path.join(BASE_DIR, "static", "exports")

# 数学/代码占位符：Unicode 私有区字符，正常笔记不会出现
_MATH_START, _MATH_END = "\uE000", "\uE001"
_CODE_START, _CODE_END = "\uE002", "\uE003"

_DISPLAY_MATH_RE = re.compile(r"(?<!\\)\$\$(.+?)\$\$", re.DOTALL)
# 行内公式收口：内容首尾不留空、闭 $ 后不跟字母数字（"$5 and $10" 这类金额不误判）
_INLINE_MATH_RE = re.compile(r"(?<!\\)\$(?!\s)((?:\\.|[^$\n])+?)(?<!\s)\$(?![0-9a-zA-Z])")
_INLINE_CODE_RE = re.compile(r"`+[^`\n]+`+")

MAX_DOCX_IMAGE_WIDTH_IN = 6.0

_PDF_CSS = """
body { font-size: 10.5pt; }
h1 { font-size: 19pt; }
h2 { font-size: 16pt; }
h3 { font-size: 13.5pt; }
h4, h5, h6 { font-size: 12pt; }
table { border-collapse: collapse; width: 100%; }
th, td { border: 1px solid #999999; padding: 4px 6px; }
th { background-color: #f2f2f2; }
img { max-width: 100%; }
code { font-family: monospace; background-color: #f5f5f5; }
pre { background-color: #f5f5f5; padding: 8px; }
blockquote { color: #555555; margin-left: 12px; padding-left: 10px; border-left: 3px solid #dddddd; }
"""


def safe_filename(title: str, fallback: str = "note") -> str:
    title = (title or "").strip() or fallback
    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", title).strip(" ._")
    return (cleaned or fallback)[:80]


# ---------------------------------------------------------------------------
# 数学公式预处理
# ---------------------------------------------------------------------------

def _split_fences(md_text: str):
    """按行切出围栏代码块，返回 [(is_fence, text), ...]，围栏内原文透传。"""
    segments, fence_buf, in_fence, marker = [], [], False, ""
    for line in md_text.split("\n"):
        stripped = line.lstrip()
        if not in_fence and (stripped.startswith("```") or stripped.startswith("~~~")):
            in_fence, marker, fence_buf = True, stripped[:3], [line]
        elif in_fence and stripped.startswith(marker):
            fence_buf.append(line)
            segments.append((True, "\n".join(fence_buf)))
            in_fence, fence_buf = False, []
        elif in_fence:
            fence_buf.append(line)
        else:
            segments.append((False, line))
    if fence_buf:
        segments.append((True, "\n".join(fence_buf)))
    return segments


def _substitute_chunk(chunk: str, math_items):
    """对一段非围栏文本：先保护行内代码，再挖显示/行内公式成占位符。"""
    code_map = []

    def _stash_code(m):
        code_map.append(m.group(0))
        return f"{_CODE_START}{len(code_map) - 1}{_CODE_END}"

    chunk = _INLINE_CODE_RE.sub(_stash_code, chunk)

    def _stash_display(m):
        math_items.append({"latex": m.group(1).strip(), "display": True})
        return f"\n\n{_MATH_START}{len(math_items) - 1}{_MATH_END}\n\n"

    def _stash_inline(m):
        math_items.append({"latex": m.group(1), "display": False})
        return f"{_MATH_START}{len(math_items) - 1}{_MATH_END}"

    chunk = _DISPLAY_MATH_RE.sub(_stash_display, chunk)
    chunk = _INLINE_MATH_RE.sub(_stash_inline, chunk)

    return re.sub(f"{re.escape(_CODE_START)}(\\d+){re.escape(_CODE_END)}",
                  lambda m: code_map[int(m.group(1))], chunk)


def extract_math(markdown: str):
    """返回 (替换后的 markdown, math_items)；占位符 \uE000{i}\uE001 留待渲染端还原。"""
    math_items = []
    out = []
    buf = []

    def flush():
        if buf:
            out.append(_substitute_chunk("\n".join(buf), math_items))
            buf.clear()

    for is_fence, text in _split_fences(markdown):
        if is_fence:
            flush()
            out.append(text)
        else:
            buf.append(text)
    flush()
    return "\n".join(out), math_items


# ---------------------------------------------------------------------------
# 数学公式渲染
# ---------------------------------------------------------------------------

def render_math_png(latex: str, display: bool = False, dpi: int = 300):
    """mathtext 渲 PNG；不认识的 LaTeX 语法返回 None（由调用方降级）。"""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.mathtext as mathtext

        expr = latex.strip()
        if not expr.startswith("$"):
            expr = f"${expr}$"
        buf = io.BytesIO()
        mathtext.math_to_image(expr, buf, dpi=dpi, format="png", color="#1f2328")
        data = buf.getvalue()
        return data if len(data) > 0 else None
    except Exception:
        return None


def latex_to_omml_element(latex: str, display: bool = False):
    """LaTeX → OMML 原生公式元素（Word 里可编辑）；失败返回 None。"""
    try:
        import latex2mathml.converter
        import mathml2omml
        from docx.oxml import parse_xml

        mathml = latex2mathml.converter.convert(latex.strip())
        omml = mathml2omml.convert(mathml)
        m_ns = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
        xml = f"<m:oMathPara {m_ns}>{omml}</m:oMathPara>" if display else omml.replace("<m:oMath>", f"<m:oMath {m_ns}>", 1)
        return parse_xml(xml)
    except Exception:
        return None


# ---------------------------------------------------------------------------
# 图片定位
# ---------------------------------------------------------------------------

def resolve_image_bytes(src: str):
    """markdown 图片目标 → 图片 bytes；解析失败返回 None。"""
    src = (src or "").strip()
    if not src:
        return None
    if src.startswith("data:"):
        try:
            return base64.b64decode(src.split(",", 1)[1])
        except Exception:
            return None
    if src.startswith(("http://", "https://")):
        try:
            resp = httpx.get(src, timeout=15, follow_redirects=True)
            resp.raise_for_status()
            return resp.content
        except Exception:
            return None
    # markdown-it 会把链接目标做 URL 编码（反斜杠 → %5C），本地路径先还原
    raw = unquote(src)
    if raw.lower().startswith("file:///"):
        raw = raw[len("file:///"):]
    candidates = []
    if raw.startswith("/"):
        candidates.append(os.path.join(BASE_DIR, raw.lstrip("/")))
    candidates += [os.path.abspath(raw), os.path.join(BASE_DIR, raw)]
    for path in candidates:
        try:
            if os.path.isfile(path):
                with open(path, "rb") as f:
                    return f.read()
        except Exception:
            continue
    return None


def _guess_mime(src: str, data: bytes) -> str:
    target = urlparse(src).path if "://" in src else src
    mime = mimetypes.guess_type(target)[0]
    if mime:
        return mime
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    return "image/png"


def embed_local_images(markdown: str) -> str:
    """把 markdown 里所有非 data-URI 的图片目标转成 base64 data-URI，供 Story 离线渲染。"""
    pattern = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)\)")

    def repl(m):
        alt, src = m.group(1), m.group(2)
        if src.startswith("data:"):
            return m.group(0)
        data = resolve_image_bytes(src)
        if not data:
            return m.group(0)
        b64 = base64.b64encode(data).decode("ascii")
        return f"![{alt}](data:{_guess_mime(src, data)};base64,{b64})"

    return pattern.sub(repl, markdown)


# ---------------------------------------------------------------------------
# PDF（markdown-it-py → HTML → fitz.Story 分页）
# ---------------------------------------------------------------------------

_MD_PARSER = None


def _md_parser():
    global _MD_PARSER
    if _MD_PARSER is None:
        from markdown_it import MarkdownIt

        _MD_PARSER = MarkdownIt("commonmark").enable(["table", "strikethrough"])
    return _MD_PARSER


def _render_math_placeholders_for_pdf(md_text: str, math_items) -> str:
    """把公式占位符替换为带显式尺寸的 raw HTML <img>。

    尺寸必须显式给：mathtext 按 dpi=300 出图，Story 会按原始像素渲染，
    行内公式会大到离谱。按 72/300 换算回排版点数。
    """
    from PIL import Image as PILImage

    for i, item in enumerate(math_items):
        placeholder = f"{_MATH_START}{i}{_MATH_END}"
        png = render_math_png(item["latex"], item["display"])
        if png is None:
            # 降级：代码体保留源码（markdown 代码 span 不吃反斜杠）
            md_text = md_text.replace(placeholder, f"`{item['latex']}`")
            continue
        with PILImage.open(io.BytesIO(png)) as im:
            w_px, h_px = im.size
        h_pt = h_px * 72.0 / 300.0
        w_pt = w_px * 72.0 / 300.0
        b64 = base64.b64encode(png).decode("ascii")
        if item["display"]:
            h_pt = min(max(h_pt, 14.0), 160.0)
            if w_pt > 460:  # A4 可用宽约 523pt，超宽公式按宽缩
                h_pt *= 460.0 / w_pt
            md_text = md_text.replace(placeholder, (
                f'<p style="text-align:center">'
                f'<img src="data:image/png;base64,{b64}" style="height:{h_pt:.1f}pt" /></p>'
            ))
        else:
            h_pt = min(max(h_pt, 7.0), 20.0)
            md_text = md_text.replace(placeholder, (
                f'<img src="data:image/png;base64,{b64}" '
                f'style="height:{h_pt:.1f}pt; vertical-align:-30%" />'
            ))
    return md_text


def export_pdf(markdown: str, title: str, out_dir: str = None) -> str:
    import fitz

    md_text, math_items = extract_math(markdown.strip())
    if math_items:
        md_text = _render_math_placeholders_for_pdf(md_text, math_items)
    html = _md_parser().render(embed_local_images(md_text))

    out_dir = out_dir or EXPORT_DIR
    os.makedirs(out_dir, exist_ok=True)
    save_path = os.path.join(out_dir, f"{uuid.uuid4().hex}_{safe_filename(title)}.pdf")

    story = fitz.Story(html=html, user_css=_PDF_CSS)
    writer = fitz.DocumentWriter(save_path)
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (36, 36, -36, -36)
    more, pages = True, 0
    try:
        while more and pages < 1000:
            dev = writer.begin_page(mediabox)
            more, _filled = story.place(where)
            story.draw(dev)
            writer.end_page()
            pages += 1
    finally:
        writer.close()
    return save_path


# ---------------------------------------------------------------------------
# Word（markdown-it-py token 流 → python-docx）
# ---------------------------------------------------------------------------

def _docx_add_hyperlink(paragraph, text: str, url: str):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


class _DocxRenderer:
    def __init__(self, doc, math_items):
        self.doc = doc
        self.math_items = math_items
        self.list_stack = []      # 每层 {"ordered": bool, "count": int}
        self.quote_depth = 0
        self.link_stack = []

    # ---- 入口 ----
    def render(self, tokens):
        i = 0
        while i < len(tokens):
            t = tokens[i]
            tt = t.type
            if tt == "heading_open":
                level = min(int(t.tag[1]), 9)
                p = self._add_heading(level)
                self.render_inline(tokens[i + 1], p)
                i += 3  # heading_open, inline, heading_close
                continue
            if tt == "paragraph_open":
                p = self._add_paragraph()
                self.render_inline(tokens[i + 1], p)
                i += 3  # paragraph_open, inline, paragraph_close
                continue
            if tt in ("bullet_list_open", "ordered_list_open"):
                self.list_stack.append({"ordered": tt == "ordered_list_open", "count": 0})
                i += 1
                continue
            if tt in ("bullet_list_close", "ordered_list_close"):
                if self.list_stack:
                    self.list_stack.pop()
                i += 1
                continue
            if tt == "list_item_open":
                if self.list_stack:
                    self.list_stack[-1]["count"] += 1
                i += 1
                continue
            if tt == "blockquote_open":
                self.quote_depth += 1
                i += 1
                continue
            if tt == "blockquote_close":
                self.quote_depth = max(0, self.quote_depth - 1)
                i += 1
                continue
            if tt in ("fence", "code_block"):
                self._add_code_block(t.content)
                i += 1
                continue
            if tt == "hr":
                self._add_hr()
                i += 1
                continue
            if tt == "table_open":
                i = self._render_table(tokens, i)
                continue
            if tt == "math_block":  # 无 math 插件理论上不出现，保险
                self._add_math_paragraph(t.content.strip())
                i += 1
                continue
            # html_block / 其他 *_open、*_close 直接跳过
            i += 1
        return self.doc

    # ---- 块级 ----
    def _add_heading(self, level: int):
        try:
            return self.doc.add_heading("", level)
        except KeyError:
            return self.doc.add_paragraph()

    def _add_paragraph(self):
        from docx.shared import Inches

        p = self.doc.add_paragraph()
        depth = len(self.list_stack)
        if depth > 0:
            layer = self.list_stack[-1]
            p.paragraph_format.left_indent = Inches(min(depth, 5) * 0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            marker = f"{layer['count']}. " if layer["ordered"] else ("• " if depth <= 3 else "– ")
            p.add_run(marker)
        if self.quote_depth > 0:
            try:
                p.style = self.doc.styles["Quote"]
            except KeyError:
                pass
        return p

    def _add_code_block(self, content: str):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls, qn
        from docx.shared import Pt

        p = self.doc.add_paragraph()
        p._p.get_or_add_pPr().append(parse_xml(r'<w:shd %s w:val="clear" w:fill="F5F5F5"/>' % nsdecls("w")))
        lines = content.rstrip("\n").split("\n")
        for k, line in enumerate(lines):
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Consolas")
            if k < len(lines) - 1:
                run.add_break()

    def _add_hr(self):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        p = self.doc.add_paragraph()
        p._p.get_or_add_pPr().append(parse_xml(
            r'<w:pBdr %s><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/></w:pBdr>' % nsdecls("w")
        ))

    def _add_math_paragraph(self, latex: str):
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        element = latex_to_omml_element(latex, display=True)
        if element is not None:
            p._p.append(element)
            return
        png = render_math_png(latex, display=True)
        if png:
            self._add_picture_run(p, png)
        else:
            p.add_run(latex)

    def _render_table(self, tokens, i: int):
        rows, current_row, cell_children = [], None, None
        j = i + 1
        while j < len(tokens) and tokens[j].type != "table_close":
            t = tokens[j]
            if t.type == "tr_open":
                current_row = []
            elif t.type in ("th_open", "td_open"):
                cell_children = []
            elif t.type in ("th_close", "td_close"):
                current_row.append(cell_children or [])
                cell_children = None
            elif t.type == "tr_close":
                rows.append(current_row)
            elif t.type == "inline" and cell_children is not None:
                cell_children.extend(t.children or [])
            j += 1
        if rows:
            cols = max(len(r) for r in rows)
            table = self.doc.add_table(rows=len(rows), cols=cols)
            try:
                table.style = self.doc.styles["Table Grid"]
            except KeyError:
                pass
            for r, row in enumerate(rows):
                for c, children in enumerate(row):
                    cell = table.cell(r, c)
                    self.render_inline_children(children, cell.paragraphs[0])
        return j + 1

    # ---- 行内 ----
    def render_inline(self, inline_token, paragraph):
        self.render_inline_children(inline_token.children or [], paragraph)

    def render_inline_children(self, children, paragraph):
        fmt = {"bold": False, "italic": False, "strike": False}
        for token in children:
            tt = token.type
            if tt == "text":
                self._emit_text(paragraph, token.content, fmt)
            elif tt == "code_inline":
                from docx.shared import Pt

                run = paragraph.add_run(token.content)
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            elif tt in ("strong_open", "strong_close"):
                fmt["bold"] = tt == "strong_open"
            elif tt in ("em_open", "em_close"):
                fmt["italic"] = tt == "em_open"
            elif tt in ("s_open", "s_close"):
                fmt["strike"] = tt == "s_open"
            elif tt in ("softbreak", "hardbreak"):
                paragraph.add_run().add_break()
            elif tt == "image":
                self._add_image(paragraph, token)
            elif tt == "link_open":
                self.link_stack.append((token.attrs or {}).get("href", ""))
            elif tt == "link_close":
                if self.link_stack:
                    self.link_stack.pop()
            # math_inline / html_inline 等忽略
        # 纯公式段（占位符独占一段）居中
        if len(children) == 1 and children[0].type == "text":
            m = re.fullmatch(f"{re.escape(_MATH_START)}(\\d+){re.escape(_MATH_END)}", children[0].content)
            if m and self.math_items[int(m.group(1))].get("display"):
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _emit_text(self, paragraph, text: str, fmt):
        pattern = re.compile(f"{re.escape(_MATH_START)}(\\d+){re.escape(_MATH_END)}")
        pieces = pattern.split(text)
        for k, piece in enumerate(pieces):
            if k % 2 == 0:
                if not piece:
                    continue
                if self.link_stack:
                    _docx_add_hyperlink(paragraph, piece, self.link_stack[-1])
                else:
                    run = paragraph.add_run(piece)
                    run.bold = fmt["bold"]
                    run.italic = fmt["italic"]
                    run.font.strike = fmt["strike"]
            else:
                item = self.math_items[int(piece)]
                element = latex_to_omml_element(item["latex"], item["display"])
                if element is not None:
                    paragraph._p.append(element)
                    continue
                png = render_math_png(item["latex"], item["display"])
                if png:
                    run = self._add_picture_run(paragraph, png)
                    run.bold = fmt["bold"]
                else:
                    run = paragraph.add_run(item["latex"])
                    run.italic = True

    def _add_image(self, paragraph, token):
        src = (token.attrs or {}).get("src", "")
        data = resolve_image_bytes(src)
        if data:
            self._add_picture_run(paragraph, data)
        else:
            alt = token.children[0].content if token.children else ""
            paragraph.add_run(f"[图片缺失: {alt or src}]")

    def _add_picture_run(self, paragraph, data: bytes):
        from docx.shared import Inches
        from PIL import Image as PILImage

        run = paragraph.add_run()
        try:
            with PILImage.open(io.BytesIO(data)) as im:
                width_in = im.size[0] / 96.0
        except Exception:
            width_in = MAX_DOCX_IMAGE_WIDTH_IN
        width_in = min(width_in, MAX_DOCX_IMAGE_WIDTH_IN)
        run.add_picture(io.BytesIO(data), width=Inches(width_in))
        return run


def export_docx(markdown: str, title: str, out_dir: str = None) -> str:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    md_text, math_items = extract_math(markdown.strip())
    tokens = _md_parser().parse(md_text)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")

    doc.core_properties.title = title

    _DocxRenderer(doc, math_items).render(tokens)

    out_dir = out_dir or EXPORT_DIR
    os.makedirs(out_dir, exist_ok=True)
    save_path = os.path.join(out_dir, f"{uuid.uuid4().hex}_{safe_filename(title)}.docx")
    doc.save(save_path)
    return save_path
